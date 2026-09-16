"""Extracción de texto y de señales de maquetación desde PDF, DOCX y texto plano.

Lo que un ATS "ve" es exactamente lo que un extractor como este consigue leer:
por eso el mismo paso que saca el texto recoge también las señales de formato
(tablas, columnas, imágenes, texto oculto) que luego audita `rules.py`.
"""

from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Tuple

from . import config

_MULTISPACE_GAP = re.compile(r"\S {3,}\S")


def extract(data: bytes, filename: str) -> Dict[str, Any]:
    """Extrae texto y metadatos de maquetación de un fichero de CV.

    Devuelve siempre un dict con las mismas claves; si la extracción falla,
    `extraction_ok` es False y `notes` explica el motivo.
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    doc: Dict[str, Any] = {
        "filename": filename,
        "ext": ext,
        "pages": 0,
        "text": "",
        "chars": 0,
        "words": 0,
        "tables": 0,
        "images": 0,
        "hidden_text_runs": 0,
        "hidden_text_samples": [],
        "columns_suspected": False,
        "has_header_footer": False,
        "extraction_ok": False,
        "notes": [],
    }

    try:
        if ext == "pdf":
            _read_pdf(data, doc)
        elif ext == "docx":
            _read_docx(data, doc)
        elif ext in ("txt", "md"):
            doc["text"] = data.decode("utf-8", errors="replace")
            doc["pages"] = 1
            doc["notes"].append(
                "Fichero de texto plano: no se pueden auditar tablas, columnas ni texto oculto."
            )
        else:
            doc["notes"].append(
                f"Extensión '.{ext}' no soportada. Usa {', '.join(config.SUPPORTED_EXTENSIONS)}."
            )
            return doc
    except Exception as exc:  # noqa: BLE001 - cualquier fallo del parser degrada igual
        doc["notes"].append(f"Error al leer el fichero: {type(exc).__name__}: {exc}")
        return doc

    doc["text"] = _normalize_whitespace(doc["text"])
    doc["chars"] = len(doc["text"])
    doc["words"] = len(doc["text"].split())
    doc["columns_suspected"] = _looks_multicolumn(doc["text"])
    doc["extraction_ok"] = doc["chars"] > 0
    if not doc["extraction_ok"]:
        doc["notes"].append(
            "No se extrajo ningún carácter: el documento es probablemente una imagen escaneada."
        )
    return doc


# --- PDF ------------------------------------------------------------------


def _read_pdf(data: bytes, doc: Dict[str, Any]) -> None:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:  # noqa: BLE001
            doc["notes"].append(
                "El PDF está protegido con contraseña; muchos ATS lo rechazan directamente."
            )
            return

    chunks: List[str] = []
    images = 0
    for page in reader.pages:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - una página ilegible no invalida el resto
            doc["notes"].append("Una página del PDF no se pudo extraer.")
        try:
            images += len(page.images)
        except Exception:  # noqa: BLE001 - pypdf no siempre puede enumerar recursos
            pass

    doc["pages"] = len(reader.pages)
    doc["images"] = images
    doc["text"] = "\n".join(chunks)


# --- DOCX -----------------------------------------------------------------


def _read_docx(data: bytes, doc: Dict[str, Any]) -> None:
    import docx
    from docx.shared import RGBColor

    document = docx.Document(io.BytesIO(data))
    parts: List[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    # El texto dentro de tablas se extrae, pero se contabiliza aparte: es la
    # causa número uno de CVs desordenados al pasar por un parser ATS.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    hidden, samples = _docx_hidden_runs(document, RGBColor)
    doc["tables"] = len(document.tables)
    doc["images"] = len(document.inline_shapes)
    doc["hidden_text_runs"] = hidden
    doc["hidden_text_samples"] = samples
    doc["has_header_footer"] = _docx_has_header_footer(document)
    doc["pages"] = max(1, round(sum(len(p) for p in parts) / 2800) or 1)
    doc["text"] = "\n".join(parts)


def _docx_hidden_runs(document: Any, rgb_cls: Any) -> Tuple[int, List[str]]:
    """Cuenta runs invisibles: marcados como ocultos, en blanco o microscópicos.

    Es la técnica clásica de "keyword stuffing" invisible y también el vehículo
    habitual de una inyección de prompt dentro de un CV.
    """
    white = rgb_cls(0xFF, 0xFF, 0xFF)
    hidden = 0
    samples: List[str] = []
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            text = run.text.strip()
            if not text:
                continue
            font = run.font
            is_hidden = bool(getattr(font, "hidden", False))
            try:
                is_white = font.color is not None and font.color.rgb == white
            except Exception:  # noqa: BLE001 - color por tema, sin rgb accesible
                is_white = False
            is_tiny = font.size is not None and font.size.pt < 4
            if is_hidden or is_white or is_tiny:
                hidden += 1
                if len(samples) < 5:
                    samples.append(text[:160])
    return hidden, samples


def _docx_has_header_footer(document: Any) -> bool:
    for section in document.sections:
        for container in (section.header, section.footer):
            if any(p.text.strip() for p in container.paragraphs):
                return True
    return False


# --- Utilidades -----------------------------------------------------------


def _normalize_whitespace(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\t\x0b\x0c]", " ", text)
    text = re.sub(r" {4,}", "   ", text)  # conserva el hueco, acota el ruido
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _looks_multicolumn(text: str) -> bool:
    """Heurística de doble columna: muchas líneas con un hueco central amplio."""
    lines = [ln for ln in text.split("\n") if len(ln.strip()) > 25]
    if len(lines) < 12:
        return False
    gapped = sum(1 for ln in lines if _MULTISPACE_GAP.search(ln))
    return gapped / len(lines) > 0.30
